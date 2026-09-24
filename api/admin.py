from fastapi import APIRouter, Depends, HTTPException, Query
from typing import Optional, List, Dict
import requests
import os
import json
import logging
logger = logging.getLogger(__name__)
from pydantic import BaseModel, Field

class AdminMutationRequest(BaseModel):
    reason: Optional[str] = Field(None, max_length=500)

from api.auth.entitlements import get_current_user, get_user_role, get_user_plan_and_status, audit_log, require_admin

admin_router = APIRouter(prefix="/api/admin", tags=["admin"])


@admin_router.get("/me")
def get_me(user: dict = Depends(require_admin)):
    return {
        "authenticated": True,
        "user_id": user.get("sub"),
        "role": "admin"
    }


@admin_router.get("/overview")
def get_overview(user: dict = Depends(require_admin)):
    import logging
    logger = logging.getLogger(__name__)

    supabase_url = os.environ.get('SUPABASE_URL', '').rstrip('/')
    supabase_key = os.environ.get('SUPABASE_SECRET_KEY', '')
    if not supabase_url or not supabase_key:
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    headers = {
        "apikey": supabase_key,
        "Authorization": f"Bearer {supabase_key}"
    }
    count_headers = {
        **headers,
        "Prefer": "count=exact"
    }

    def get_postgrest_count(table: str, filters: dict, select_col: str, label: str) -> int:
        req_params = {**filters, "limit": "1", "select": select_col}

        resp = requests.get(f"{supabase_url}/rest/v1/{table}", params=req_params, headers=count_headers, timeout=5.0)

        if resp.status_code not in (200, 206):
            logger.error(f"Overview: {label} returned HTTP {resp.status_code}")
            # Safe diagnostic detail
            raise HTTPException(status_code=502, detail=f"Overview metric failed: {label} status={resp.status_code} content_range_present={'Content-Range' in resp.headers}")

        cr = resp.headers.get("Content-Range", "")
        if "/" not in cr:
            logger.error(f"Overview: {label} missing Content-Range header")
            raise HTTPException(status_code=502, detail=f"Overview metric failed: {label} status={resp.status_code} content_range_present=false")

        try:
            return int(cr.split("/")[-1])
        except (ValueError, IndexError):
            logger.error(f"Overview: {label} unparseable Content-Range: {cr}")
            raise HTTPException(status_code=502, detail=f"Overview metric failed: {label} status={resp.status_code} content_range_present=true")

    try:
        import datetime
        from concurrent.futures import ThreadPoolExecutor

        today = datetime.datetime.now(datetime.timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
        monday = today - datetime.timedelta(days=today.weekday())

        def fetch_users():
            resp = requests.get(f"{supabase_url}/auth/v1/admin/users", headers=headers, timeout=5.0)
            if resp.status_code != 200:
                logger.error(f"Overview: Auth users returned HTTP {resp.status_code}")
                raise HTTPException(status_code=502, detail=f"Overview metric failed: auth users status={resp.status_code}")
            data = resp.json()
            t = data.get("total")
            return t if t is not None else len(data.get("users", []))

        with ThreadPoolExecutor(max_workers=6) as executor:
            fut_users = executor.submit(fetch_users)
            fut_pro = executor.submit(get_postgrest_count, "user_plans", {"plan": "eq.professional"}, "user_id", "professional users")
            fut_susp = executor.submit(get_postgrest_count, "user_plans", {"status": "eq.suspended"}, "user_id", "suspended users")
            fut_tot_scans = executor.submit(get_postgrest_count, "scans", {}, "id", "total scans")
            fut_today = executor.submit(get_postgrest_count, "scans", {"created_at": f"gte.{today.isoformat()}"}, "id", "scans today")
            fut_week = executor.submit(get_postgrest_count, "scans", {"created_at": f"gte.{monday.isoformat()}"}, "id", "scans this week")

            total_users = fut_users.result()
            professional_users = fut_pro.result()
            suspended_users = fut_susp.result()
            total_scans = fut_tot_scans.result()
            scans_today = fut_today.result()
            scans_this_week = fut_week.result()

        free_users = total_users - professional_users
        active_users = total_users - suspended_users

        return {
            "total_users": total_users,
            "free_users": free_users,
            "professional_users": professional_users,
            "active_users": active_users,
            "suspended_users": suspended_users,
            "total_scans": total_scans,
            "scans_today": scans_today,
            "scans_this_week": scans_this_week
        }
    except HTTPException:
        raise
    except requests.exceptions.RequestException as e:
        logger.error(f"Overview upstream failure: {e}")
        raise HTTPException(status_code=502, detail="Upstream API failure")
    except Exception as e:
        logger.error(f"Overview unexpected error: {e}")
        raise HTTPException(status_code=500, detail="Internal server error")


def _fetch_users_page(url_override: str = None) -> dict:
    url = url_override or f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/auth/v1/admin/users"
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
        "Content-Type": "application/json"
    }
    return requests.get(url, headers=headers, timeout=5.0)

def _get_mapped_users(users_data: list, plans_map: dict, roles_map: dict, search: str = None) -> list:
    safe_users = []
    for u in users_data:
        uid = u.get("id")
        if search:
            s = search.lower()
            email = u.get("email", "").lower()
            if s not in email and s not in uid.lower():
                continue

        plan_info = plans_map.get(uid, {})
        safe_users.append({
            "user_id": uid,
            "email": u.get("email"),
            "name": u.get("user_metadata", {}).get("full_name", ""),
            "phone": u.get("phone"),
            "phone_confirmed_at": u.get("phone_confirmed_at"),
            "role": roles_map.get(uid, "user"),
            "plan": plan_info.get("plan", "free"),
            "status": plan_info.get("status", "active"),
            "created_at": u.get("created_at")
        })
    return safe_users

def _fetch_roles_and_plans():
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
        "Content-Type": "application/json"
    }
    def fetch_plans():
        return requests.get(f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/user_plans?select=user_id,plan,status", headers=headers, timeout=5.0)
    def fetch_roles():
        return requests.get(f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/user_roles?select=user_id,role", headers=headers, timeout=5.0)

    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=2) as executor:
        fut_plans = executor.submit(fetch_plans)
        fut_roles = executor.submit(fetch_roles)
        plans_resp = fut_plans.result()
        roles_resp = fut_roles.result()

    plans_map = {}
    if plans_resp.status_code == 200:
        for row in plans_resp.json():
            plans_map[row.get("user_id")] = row

    roles_map = {}
    if roles_resp.status_code == 200:
        for row in roles_resp.json():
            roles_map[row.get("user_id")] = row.get("role")

    return plans_map, roles_map


def _fetch_all_roles_and_plans():
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
        "Content-Type": "application/json"
    }
    base_url = os.environ.get('SUPABASE_URL', '').rstrip('/')

    def fetch_paginated(endpoint, select_fields):
        all_records = []
        limit = 1000
        offset = 0
        max_requests = 50
        requests_made = 0

        while requests_made <= max_requests:
            url = f"{base_url}/rest/v1/{endpoint}?select={select_fields}&limit={limit}&offset={offset}"
            resp = requests.get(url, headers=headers, timeout=5.0)
            if resp.status_code != 200:
                logger.error(f"Error fetching {endpoint} offset {offset}: {resp.status_code}")
                raise HTTPException(status_code=500, detail="Error fetching users")

            data = resp.json()
            if not data:
                break

            all_records.extend(data)
            if len(data) < limit:
                break

            offset += limit
            requests_made += 1

        if requests_made > max_requests:
            logger.error(f"Export exceeded max size for {endpoint}")
            raise HTTPException(status_code=500, detail="Export exceeds maximum supported size")

        return all_records

    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=2) as executor:
        fut_plans = executor.submit(fetch_paginated, "user_plans", "user_id,plan,status")
        fut_roles = executor.submit(fetch_paginated, "user_roles", "user_id,role")
        plans_data = fut_plans.result()
        roles_data = fut_roles.result()

    plans_map = {row.get("user_id"): row for row in plans_data if isinstance(row, dict)}
    roles_map = {row.get("user_id"): row.get("role") for row in roles_data if isinstance(row, dict)}

    return plans_map, roles_map

def _fetch_all_users(search: Optional[str] = None) -> list:
    if not os.environ.get('SUPABASE_URL') or not os.environ.get('SUPABASE_SECRET_KEY'):
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    try:
        plans_map, roles_map = _fetch_all_roles_and_plans()
        all_users = []
        page = 1
        per_page = 1000
        max_pages = 50

        while page <= max_pages:
            url = f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/auth/v1/admin/users?page={page}&per_page={per_page}"
            resp = _fetch_users_page(url)
            if resp.status_code != 200:
                logger.error(f"Error fetching users page {page}: {resp.status_code}")
                raise HTTPException(status_code=500, detail="Error fetching users")

            users_page = resp.json().get("users", [])
            if not users_page:
                break

            all_users.extend(users_page)
            if len(users_page) < per_page:
                break

            page += 1

        if page > max_pages:
            logger.error("Export exceeded maximum supported size.")
            raise HTTPException(status_code=500, detail="Export exceeds maximum supported size")

        return _get_mapped_users(all_users, plans_map, roles_map, search)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching all users: {e}")
        raise HTTPException(status_code=500, detail="Error fetching users")

@admin_router.get("/users")
def get_users(limit: int = Query(50), offset: int = Query(0), search: Optional[str] = Query(None), user: dict = Depends(require_admin)):
    if not os.environ.get('SUPABASE_URL') or not os.environ.get('SUPABASE_SECRET_KEY'):
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    try:
        from concurrent.futures import ThreadPoolExecutor
        def fetch_auth_users():
            # Original bounded retrieval (defaults to Supabase page 1 limits, 50 users)
            return _fetch_users_page()

        def fetch_plans():
            headers = {
                "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
                "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
                "Content-Type": "application/json"
            }
            return requests.get(f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/user_plans?select=user_id,plan,status", headers=headers, timeout=5.0)

        def fetch_roles():
            headers = {
                "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
                "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
                "Content-Type": "application/json"
            }
            return requests.get(f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/user_roles?select=user_id,role", headers=headers, timeout=5.0)

        with ThreadPoolExecutor(max_workers=3) as executor:
            fut_users = executor.submit(fetch_auth_users)
            fut_plans = executor.submit(fetch_plans)
            fut_roles = executor.submit(fetch_roles)

            resp = fut_users.result()
            plans_resp = fut_plans.result()
            roles_resp = fut_roles.result()

        if resp.status_code == 200:
            users_data = resp.json().get("users", [])

            plans_map = {}
            if plans_resp.status_code == 200:
                for row in plans_resp.json():
                    plans_map[row.get("user_id")] = row

            roles_map = {}
            if roles_resp.status_code == 200:
                for row in roles_resp.json():
                    roles_map[row.get("user_id")] = row.get("role")

            safe_users = _get_mapped_users(users_data, plans_map, roles_map, search)
            return safe_users[offset:offset+limit]
    except Exception as e:
        logger.error(f"Error fetching users: {e}")
        raise HTTPException(status_code=500, detail="Error fetching users")
    return []

from fastapi.responses import Response
import csv
import io
import datetime

def safe_export_val(val) -> str:
    if val is None:
        return ""
    v = str(val).strip()
    if v.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{v}"
    return v

@admin_router.get("/users/export")
def export_users(format: str = Query("csv"), search: Optional[str] = Query(None), role: str = Query("all"), plan: str = Query("all"), status: str = Query("all"), user: dict = Depends(require_admin)):
    all_users = _fetch_all_users(search)

    filtered_users = []
    for u in all_users:
        if role != "all" and u.get("role", "user") != role:
            continue
        if plan != "all" and u.get("plan", "free") != plan:
            continue
        if status != "all" and u.get("status", "active") != status:
            continue
        filtered_users.append(u)

    headers = [
        "User ID", "Name", "Email", "Phone", "Phone Verification Status",
        "Role", "Plan", "Status", "Created At"
    ]

    rows = []
    for u in filtered_users:
        phone = u.get("phone")
        phone_confirmed_at = u.get("phone_confirmed_at")

        verification = "Not provided"
        if phone:
            if phone_confirmed_at:
                verification = "Verified"
            else:
                verification = "Unverified"

        rows.append([
            safe_export_val(u.get("user_id")),
            safe_export_val(u.get("name")),
            safe_export_val(u.get("email")),
            safe_export_val(phone if phone else "Not provided"),
            safe_export_val(verification),
            safe_export_val(u.get("role", "user")),
            safe_export_val(u.get("plan", "free").capitalize() if u.get("plan") else "Free"),
            safe_export_val(u.get("status", "active")),
            safe_export_val(u.get("created_at"))
        ])

    timestamp = datetime.datetime.now().strftime("%Y-%m-%d")

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)

        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=urlscanonline-users-{timestamp}.csv"}
        )
    elif format == "xlsx":
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise HTTPException(status_code=500, detail="XLSX export not supported")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Users"

        ws.append(headers)
        for row in rows:
            ws.append(row)

        for i, col in enumerate(ws.columns, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = 20
            ws[f"{col_letter}1"].font = openpyxl.styles.Font(bold=True)

        output = io.BytesIO()
        wb.save(output)

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=urlscanonline-users-{timestamp}.xlsx"}
        )
    else:
        raise HTTPException(status_code=400, detail="Invalid format")

@admin_router.get("/users/{user_id}")
def get_user_detail(user_id: str, user: dict = Depends(require_admin)):
    if not os.environ.get('SUPABASE_URL') or not os.environ.get('SUPABASE_SECRET_KEY'):
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    url = f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/auth/v1/admin/users/{user_id}"
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
        "Content-Type": "application/json"
    }

    plan, status = get_user_plan_and_status(user_id)

    try:
        resp = requests.get(url, headers=headers, timeout=5.0)
        if resp.status_code == 200:
            u = resp.json()
            return {
                "user_id": u.get("id"),
                "email": u.get("email"),
                "role": get_user_role(user_id),
                "plan": plan,
                "status": status,
                "created_at": u.get("created_at")
            }
    except Exception as e:
        pass

    return {
        "user_id": user_id,
        "role": get_user_role(user_id),
        "plan": plan,
        "status": status
    }

def upsert_user_plan(user_id: str, new_plan: str, new_status: str):
    url = f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/user_plans"
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"
    }
    payload = {
        "user_id": user_id,
        "plan": new_plan,
        "status": new_status
    }
    resp = requests.post(url, json=payload, headers=headers, timeout=5.0)
    if resp.status_code not in (200, 201, 204):
        raise HTTPException(status_code=500, detail="Failed to update user_plans")

def verify_user_exists(user_id: str):
    if not os.environ.get('SUPABASE_URL') or not os.environ.get('SUPABASE_SECRET_KEY'):
        return
    url = f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/auth/v1/admin/users/{user_id}"
    headers = {
        "apikey": os.environ.get("SUPABASE_SECRET_KEY", ""),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}"
    }
    resp = requests.get(url, headers=headers, timeout=5.0)
    if resp.status_code == 404:
        raise HTTPException(status_code=404, detail="User not found.")
    elif resp.status_code != 200:
        raise HTTPException(status_code=500, detail="Error verifying user existence.")

@admin_router.post("/users/{user_id}/grant-professional")
def grant_professional(user_id: str, payload: Optional[AdminMutationRequest] = None, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    current_role = get_user_role(user_id)
    current_plan, current_status = get_user_plan_and_status(user_id)
    before_state = {"role": current_role, "plan": current_plan, "status": current_status}

    upsert_user_plan(user_id, "professional", current_status)

    after_state = {"role": current_role, "plan": "professional", "status": current_status}
    audit_log(user.get("sub"), "grant_professional", "user", user_id, payload.reason if payload else None, before_state, after_state)
    return {"user_id": user_id, "plan": "professional", "status": current_status}

@admin_router.post("/users/{user_id}/remove-professional")
def remove_professional(user_id: str, payload: Optional[AdminMutationRequest] = None, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    current_role = get_user_role(user_id)
    current_plan, current_status = get_user_plan_and_status(user_id)
    before_state = {"role": current_role, "plan": current_plan, "status": current_status}

    upsert_user_plan(user_id, "free", current_status)

    after_state = {"role": current_role, "plan": "free", "status": current_status}
    audit_log(user.get("sub"), "remove_professional", "user", user_id, payload.reason if payload else None, before_state, after_state)
    return {"user_id": user_id, "plan": "free", "status": current_status}

@admin_router.post("/users/{user_id}/suspend")
def suspend_user(user_id: str, payload: Optional[AdminMutationRequest] = None, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    if user.get("sub") == user_id:
        raise HTTPException(status_code=400, detail="You cannot suspend your own Admin account.")

    current_role = get_user_role(user_id)
    current_plan, current_status = get_user_plan_and_status(user_id)
    before_state = {"role": current_role, "plan": current_plan, "status": current_status}

    upsert_user_plan(user_id, current_plan, "suspended")

    after_state = {"role": current_role, "plan": current_plan, "status": "suspended"}
    audit_log(user.get("sub"), "suspend_user", "user", user_id, payload.reason if payload else None, before_state, after_state)
    return {"user_id": user_id, "plan": current_plan, "status": "suspended"}

@admin_router.post("/users/{user_id}/reactivate")
def reactivate_user(user_id: str, payload: Optional[AdminMutationRequest] = None, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    current_role = get_user_role(user_id)
    current_plan, current_status = get_user_plan_and_status(user_id)
    before_state = {"role": current_role, "plan": current_plan, "status": current_status}

    upsert_user_plan(user_id, current_plan, "active")

    after_state = {"role": current_role, "plan": current_plan, "status": "active"}
    audit_log(user.get("sub"), "reactivate_user", "user", user_id, payload.reason if payload else None, before_state, after_state)
    return {"user_id": user_id, "plan": current_plan, "status": "active"}



@admin_router.get("/audit-logs")
def get_audit_logs(limit: int = Query(50), offset: int = Query(0), search: Optional[str] = Query(None), user: dict = Depends(require_admin)):
    if not os.environ.get('SUPABASE_URL') or not os.environ.get('SUPABASE_SECRET_KEY'):
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    query = f"select=*&limit={limit}&offset={offset}&order=created_at.desc"
    if search:
        if len(search) >= 8 and "-" in search:
            query += f"&or=(admin_user_id.eq.{search},resource_id.eq.{search},action.ilike.*{search}*)"
        else:
            query += f"&action=ilike.*{search}*"

    url = f"{os.environ.get('SUPABASE_URL', '').rstrip('/')}/rest/v1/audit_logs?{query}"

    headers = {
        "apikey": os.environ.get('SUPABASE_SECRET_KEY', ''),
        "Authorization": f"Bearer {os.environ.get('SUPABASE_SECRET_KEY', '')}"
    }

    try:
        resp = requests.get(url, headers=headers, timeout=5.0)
        if resp.status_code == 200:
            return resp.json()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return []


from .auth.entitlements import check_free_quota, reset_free_quota

@admin_router.get("/users/{user_id}/quota")
def get_user_quota(user_id: str, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    # Check if they are admin or pro
    from .auth.entitlements import get_user_role, get_user_plan_and_status
    role = get_user_role(user_id)
    plan, _ = get_user_plan_and_status(user_id)

    if role == "admin":
        return {"limit": "Unlimited", "used": 0, "remaining": "Unlimited", "reset_time": None}

    if plan == "professional":
        return {"limit": "Professional", "used": 0, "remaining": "Unlimited", "reset_time": None}

    # Free
    quota = check_free_quota(user_id)
    return {
        "limit": quota.get("quota_limit", 5),
        "used": quota.get("quota_used", 0),
        "remaining": quota.get("quota_remaining", 5),
        "reset_time": quota.get("reset_at")
    }

@admin_router.post("/users/{user_id}/reset-quota")
def admin_reset_quota(user_id: str, req: Optional[AdminMutationRequest] = None, user: dict = Depends(require_admin)):
    verify_user_exists(user_id)
    # Capture quota state before reset for audit
    quota_before = check_free_quota(user_id)
    before_state = {
        "quota_used": quota_before.get("quota_used", 0),
        "quota_limit": quota_before.get("quota_limit", 5),
        "quota_remaining": quota_before.get("quota_remaining", 5)
    }

    success = reset_free_quota(user_id)
    if not success:
        raise HTTPException(status_code=500, detail="Failed to reset quota in Redis")

    reason = req.reason if req else "Admin requested quota reset"
    audit_log(
        admin_user_id=user["sub"],
        action="reset_free_quota",
        resource_type="user",
        resource_id=user_id,
        reason=reason,
        before_state=before_state,
        after_state={"quota_used": 0, "quota_remaining": before_state["quota_limit"]}
    )
    return {"status": "success", "message": "Free quota reset"}


from api.scanner.compare import compare_reports

@admin_router.get("/scans/compare")
def compare_admin_scans(scan_id_1: str, scan_id_2: str, user: dict = Depends(require_admin)):
    import os
    import requests
    from fastapi import HTTPException

    supabase_url = os.environ.get('SUPABASE_URL', '').rstrip('/')
    supabase_key = os.environ.get('SUPABASE_SECRET_KEY', '')
    if not supabase_url or not supabase_key:
        raise HTTPException(status_code=500, detail="Supabase credentials not configured.")

    headers = {
        "apikey": supabase_key,
        "Authorization": f"Bearer {supabase_key}"
    }

    # Fetch scan 1
    resp1 = requests.get(f"{supabase_url}/rest/v1/scans?id=eq.{scan_id_1}&select=*", headers=headers, timeout=5.0)
    if resp1.status_code != 200 or not resp1.json():
        raise HTTPException(status_code=404, detail="Scan 1 not found.")
    scan1 = resp1.json()[0]
    if scan1.get("user_id") != user.get("sub"):
        raise HTTPException(status_code=403, detail="Access denied for Scan 1.")

    # Fetch scan 2
    resp2 = requests.get(f"{supabase_url}/rest/v1/scans?id=eq.{scan_id_2}&select=*", headers=headers, timeout=5.0)
    if resp2.status_code != 200 or not resp2.json():
        raise HTTPException(status_code=404, detail="Scan 2 not found.")
    scan2 = resp2.json()[0]
    if scan2.get("user_id") != user.get("sub"):
        raise HTTPException(status_code=403, detail="Access denied for Scan 2.")

    # Sort scans chronologically
    s1_time = scan1.get("created_at", "")
    s2_time = scan2.get("created_at", "")
    if s1_time and s2_time and s1_time > s2_time:
        scan1, scan2 = scan2, scan1

    try:
        result = compare_reports(scan1, scan2)
        return result
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
